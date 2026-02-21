import io
import json
from dataclasses import dataclass
from typing import List

import streamlit as st
from docx import Document
from openai import OpenAI


@dataclass
class ReviewResult:
    filename: str
    review_outline: str
    rewritten_chapters: str


def extract_docx_text(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def build_review_prompt(book_text: str, language: str, audience: str) -> str:
    return f"""
You are an elite literary editor, story architect, and commercial publishing strategist.

Task:
1) Review the manuscript and provide an editorial review outline.
2) Rewrite the manuscript into stronger chapters with:
   - Logical plot/argument flow
   - Sequential structure
   - Time consistency
   - Higher bestseller potential

Output MUST be valid JSON with this schema:
{{
  "review_outline": [
    {{"section": "...", "findings": ["..."], "recommendations": ["..."]}}
  ],
  "rewritten_chapters": [
    {{
      "chapter_number": 1,
      "title": "...",
      "timeline_notes": "...",
      "chapter_text": "..."
    }}
  ]
}}

Requirements:
- Keep language: {language}
- Target audience: {audience}
- If the input has weak structure, infer a better chapter sequence.
- Ensure character/event chronology is internally consistent.
- Improve hooks, chapter endings, pacing, and narrative momentum.
- Keep the core intent/theme of the original manuscript.

Manuscript:
{book_text}
"""


def review_and_rewrite(client: OpenAI, model: str, book_text: str, language: str, audience: str) -> dict:
    prompt = build_review_prompt(book_text=book_text, language=language, audience=audience)

    response = client.responses.create(
        model=model,
        input=[
            {"role": "system", "content": "You are precise, structured, and output strict JSON only."},
            {"role": "user", "content": prompt},
        ],
        temperature=0.7,
    )

    text = response.output_text.strip()
    return json.loads(text)


def create_output_docx(result: ReviewResult) -> bytes:
    doc = Document()
    doc.add_heading(f"Book Review + Rewrite: {result.filename}", level=1)

    doc.add_heading("Editorial Review Outline", level=2)
    doc.add_paragraph(result.review_outline)

    doc.add_heading("Rewritten Chapters", level=2)
    doc.add_paragraph(result.rewritten_chapters)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.read()


def main() -> None:
    st.set_page_config(page_title="Book Review & Rewrite Studio", layout="wide")
    st.title("📚 Book Review & Rewrite Studio")
    st.write(
        "Upload one or more Microsoft Word manuscripts (.docx). The app reviews each file and rewrites chapters "
        "for stronger sequencing, logical flow, timeline consistency, and bestseller potential."
    )

    api_key = st.text_input(
        "OpenAI API key",
        type="password",
        placeholder="sk-...",
        help="Enter your own API key for this session. It is not pre-filled from server secrets.",
    )
    model = st.text_input("Model", value="gpt-4.1")

    col1, col2 = st.columns(2)
    with col1:
        language = st.text_input("Output language", value="English")
    with col2:
        audience = st.text_input("Target audience", value="General adult fiction readers")

    uploaded_files = st.file_uploader(
        "Upload Word files",
        type=["docx"],
        accept_multiple_files=True,
    )

    if st.button("Review and Rewrite", type="primary"):
        if not api_key:
            st.error("Please provide an OpenAI API key.")
            st.stop()
        if not uploaded_files:
            st.error("Please upload at least one .docx file.")
            st.stop()

        client = OpenAI(api_key=api_key)
        results: List[ReviewResult] = []

        for file in uploaded_files:
            with st.spinner(f"Processing {file.name}..."):
                text = extract_docx_text(file.read())
                if not text:
                    st.warning(f"{file.name} appears empty. Skipping.")
                    continue

                try:
                    data = review_and_rewrite(
                        client=client,
                        model=model,
                        book_text=text,
                        language=language,
                        audience=audience,
                    )
                except Exception as exc:
                    st.error(f"Failed for {file.name}: {exc}")
                    continue

                outline = json.dumps(data.get("review_outline", []), indent=2, ensure_ascii=False)
                chapters = json.dumps(data.get("rewritten_chapters", []), indent=2, ensure_ascii=False)
                results.append(ReviewResult(file.name, outline, chapters))

        if not results:
            st.warning("No files were successfully processed.")
            st.stop()

        st.success(f"Processed {len(results)} file(s).")

        for result in results:
            st.subheader(result.filename)
            tab1, tab2 = st.tabs(["Review Outline", "Rewritten Chapters"])
            with tab1:
                st.code(result.review_outline, language="json")
            with tab2:
                st.code(result.rewritten_chapters, language="json")

            out_bytes = create_output_docx(result)
            st.download_button(
                label=f"Download rewritten report for {result.filename}",
                data=out_bytes,
                file_name=f"rewritten_{result.filename}",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )


if __name__ == "__main__":
    main()
